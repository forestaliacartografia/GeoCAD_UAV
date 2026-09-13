"""
v1.21.0: the relazione, written three ways, and read back.

One report object, three writers. This suite builds a real project, asks the
Elaborati panel for the report, writes it to PDF, to .docx and to .xlsx, and
then *reads each file back* -- because a file that was written is not a file
that can be opened.

 * the workbook comes back through openpyxl, sheet by sheet, and a number
   written as a number comes back as a number;
 * the PDF is opened again and its text extracted (PyPDF2 when QGIS ships it,
   otherwise the structure is checked and the extraction is skipped);
 * the .docx is unzipped, every part parsed, and -- where Microsoft Word is
   installed -- opened by Word itself, which is the only authority on whether
   a hand-written OOXML package is a Word document. That block skips with its
   reason where Word is not there.

NEEDS QGIS with widgets. Run with:

    "C:\\Program Files\\QGIS 3.40.15\\bin\\python-qgis-ltr.bat" ^
        geocad_uav\\tests\\test_documents.py
"""

import gc
import os
import sys
import tempfile
import zipfile
from xml.etree import ElementTree

sys.path.insert(0, os.path.dirname(os.path.dirname(
    os.path.dirname(os.path.abspath(__file__)))))

from qgis.core import (QgsApplication,                          # noqa: E402
                       QgsCoordinateReferenceSystem, QgsGeometry, QgsProject)

QGS = QgsApplication([], True)
QGS.initQgis()

from geocad_uav.gui import workflow as wf                       # noqa: E402
from geocad_uav.io import documents as docs                     # noqa: E402

FAILURES = []
SKIPS = []
TMP = tempfile.mkdtemp(prefix="geocad_docs_")
CRS32632 = QgsCoordinateReferenceSystem("EPSG:32632")
OX, OY = 500000.0, 5000000.0


def check(label, got, expected, tol=1e-9):
    ok = abs(float(got) - float(expected)) <= tol
    print("  [{0}] {1:<54} got={2:<16.10g} exp={3:.10g}".format(
        "ok  " if ok else "FAIL", label, float(got), float(expected)))
    if not ok:
        FAILURES.append(label)


def check_text(label, got, expected):
    ok = got == expected
    print("  [{0}] {1:<46} got={2!r:<24} exp={3!r}".format(
        "ok  " if ok else "FAIL", label, got, expected))
    if not ok:
        FAILURES.append(label)


def check_true(label, condition):
    print("  [{0}] {1}".format("ok  " if condition else "FAIL", label))
    if not condition:
        FAILURES.append(label)


def skip(label, reason):
    print("  [skip] {0}\n         motivo: {1}".format(label, reason))
    SKIPS.append((label, reason))


AREA = QgsGeometry.fromWkt(
    "POLYGON(({0} {1},{2} {1},{2} {3},{0} {3},{0} {1}))".format(
        OX, OY, OX + 300.0, OY + 240.0))

print("=" * 78)
print("Elaborati -- una relazione, tre file, tutti riletti")
print("=" * 78)

workspace = wf.Workspace(None)
state = workspace.state
scheme = workspace.context.scheme_panel
zones_panel = workspace.context.zones_panel
generate_panel = workspace.context.generate_panel
outputs = workspace.context.outputs_panel

state.set_area(AREA, CRS32632, "Podere della Fonte")
state.constraints.declare("strada", 8.0, label="Strade")
state.constraints.add_geometry("strada", QgsGeometry.fromWkt(
    "LINESTRING({0} {1},{2} {1})".format(OX + 10.0, OY + 120.0,
                                         OX + 290.0, OY + 120.0)))
state.apply_constraints()
scheme.plant_distance.setValue(10.0)
scheme.row_distance.setValue(10.0)
for name, percent in (("quercia", 70.0), ("frassino", 30.0)):
    scheme.species_key.setCurrentText(name)
    scheme.species_percent.setValue(percent)
    scheme.on_add_species()
scheme.apply_scheme()
zones_panel.band_count.setValue(2)
zones_panel.split()
plan = generate_panel.preview()
generate_panel.generate()
workspace.context.verify_panel.run()
outputs.author_edit.setText("Cap. N. M. Mancini")
print("  progetto: {0:,} piante, {1} zone".format(plan.count,
                                                  len(state.zones)))

# --------------------------------------------------------------------------
# D1 - one report, and it says what the model says
# --------------------------------------------------------------------------
print("\n== D1: una sola relazione, costruita dal modello ==")
report = outputs.document()
check_true("la relazione ha un titolo", bool(report.title))
check_text("...e porta l'autore che l'operatore ha scritto", report.author,
           "Cap. N. M. Mancini")
check_true("...e la data", bool(report.date))
headings = [b.text for b in report.blocks if b.kind == docs.BLOCK_HEADING]
print("        sezioni: {0}".format(headings))
for expected in ("VINCOLI E FASCE DI RISPETTO", "SESTO (distanze reali "
                 "sul terreno)"):
    check_true("la sezione '{0}' c'e'".format(expected),
               any(expected in text for text in headings))

tables = {block.text: block for block in report.tables()}
print("        tabelle: {0}".format(sorted(tables)))
for expected in ("Vincoli e fasce di rispetto", "Zone", "Composizione",
                 "Piante"):
    check_true("la tabella '{0}' c'e'".format(expected), expected in tables)
check("la tabella delle zone ha una riga per zona",
      len(tables["Zone"].rows), len(state.zones))
check("la tabella delle piante le ha tutte",
      len(tables["Piante"].rows), plan.count)
check_true("...ma sta solo nel foglio di calcolo",
           tables["Piante"].sheet_only)
check_true("la composizione elenca le due specie",
           len(tables["Composizione"].rows) == 2)

text = outputs.build_report()
check_true("il testo mostrato nel pannello viene dalla stessa relazione",
           "Podere della Fonte" in text and "Zone" in text)
check_true("...e non contiene l'elenco delle piante",
           "Progressivo" not in text)
print("        anteprima: {0} righe".format(len(text.split("\n"))))

# --------------------------------------------------------------------------
# D2 - the workbook, read back with openpyxl
# --------------------------------------------------------------------------
print("\n== D2: il foglio di calcolo, riletto ==")
xlsx_path = os.path.join(TMP, "relazione.xlsx")
written = outputs.write_document("xlsx", xlsx_path)
check_text("il file e' stato scritto dove chiesto", written, xlsx_path)
check_true("...e non e' vuoto", os.path.getsize(xlsx_path) > 3000)

import openpyxl                                                 # noqa: E402

book = openpyxl.load_workbook(xlsx_path)
print("        fogli: {0}".format(book.sheetnames))
check_true("openpyxl lo riapre", len(book.sheetnames) >= 4)
check_text("il primo foglio e' la relazione", book.sheetnames[0], "Relazione")
check_true("c'e' un foglio per ogni tabella",
           all(docs.sheet_name(name) in book.sheetnames for name in tables))
prose = book["Relazione"]
check_text("...col titolo in testa", prose["A1"].value, report.title)
check_true("...e l'autore", any("Mancini" in str(row[0].value or "")
                                for row in prose.iter_rows(max_row=6)))

plants_sheet = book[docs.sheet_name("Piante")]
check("il foglio delle piante ha una riga per pianta, piu' l'intestazione",
      plants_sheet.max_row, plan.count + 1)
check_text("la prima colonna si chiama ID", plants_sheet["A1"].value, "ID")
first_x = plants_sheet["F2"].value
print("        prima pianta: id={0}, X={1!r}".format(
    plants_sheet["A2"].value, first_x))
check_true("una coordinata e' tornata come numero, non come testo",
           isinstance(first_x, (int, float)))
check_true("...ed e' quella della prima pianta",
           abs(float(first_x) - round(plan.plants[0].x, 3)) < 1e-6)
check_true("il foglio e' bloccato sotto l'intestazione",
           plants_sheet.freeze_panes == "A2")
zone_sheet = book[docs.sheet_name("Zone")]
check("il foglio delle zone ha una riga per zona", zone_sheet.max_row,
      len(state.zones) + 1)
book.close()

# --------------------------------------------------------------------------
# D3 - the PDF, reopened
# --------------------------------------------------------------------------
print("\n== D3: il PDF, riaperto ==")
pdf_path = os.path.join(TMP, "relazione.pdf")
written = outputs.write_document("pdf", pdf_path)
check_text("il PDF e' stato scritto", written, pdf_path)
size = os.path.getsize(pdf_path)
print("        PDF: {0:,} byte".format(size))
check_true("...e non e' vuoto", size > 5000)
with open(pdf_path, "rb") as handle:
    blob = handle.read()
check_text("comincia come un PDF", blob[:5], b"%PDF-")
check_true("...e finisce come un PDF", b"%%EOF" in blob[-2048:])

try:
    from PyPDF2 import PdfReader                                # noqa: E402
    reader = PdfReader(pdf_path)
    pages = len(reader.pages)
    extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    print("        {0} pagine, {1:,} caratteri estratti".format(
        pages, len(extracted)))
    check_true("il PDF ha delle pagine", pages >= 1)
    check_true("...e il testo torna fuori", "RELAZIONE TECNICA" in extracted)
    check_true("...col nome del progetto", "Podere" in extracted)
    check_true("...e i vincoli", "Strade" in extracted)
except ImportError as exc:
    skip("il testo del PDF si rilegge",
         "questa installazione di QGIS non comprende PyPDF2 ({0})".format(exc))

# --------------------------------------------------------------------------
# D4 - the .docx package, part by part
# --------------------------------------------------------------------------
print("\n== D4: il pacchetto .docx, parte per parte ==")
docx_path = os.path.join(TMP, "relazione.docx")
written = outputs.write_document("docx", docx_path)
check_text("il documento Word e' stato scritto", written, docx_path)
check_true("...ed e' uno zip", zipfile.is_zipfile(docx_path))

archive = zipfile.ZipFile(docx_path)
names = archive.namelist()
print("        parti: {0}".format(names))
for part in ("[Content_Types].xml", "_rels/.rels", "word/document.xml",
             "word/_rels/document.xml.rels", "word/styles.xml",
             "docProps/core.xml"):
    check_true("la parte {0} c'e'".format(part), part in names)
check_text("[Content_Types].xml e' la prima parte dell'archivio", names[0],
           "[Content_Types].xml")

for part in names:
    try:
        ElementTree.fromstring(archive.read(part))
    except ElementTree.ParseError as exc:
        check_true("{0} e' XML valido ({1})".format(part, exc), False)
check_true("ogni parte e' XML che si lascia riaprire", True)

types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
overrides = {node.get("PartName"): node.get("ContentType")
             for node in types
             if node.tag.endswith("Override")}
check_text("il tipo del documento principale e' quello di Word",
           overrides.get("/word/document.xml"), docs.CT_DOCUMENT)
check_true("...e ogni Override punta a una parte che esiste",
           all(name.lstrip("/") in names for name in overrides))

rels = ElementTree.fromstring(archive.read("_rels/.rels"))
targets = {node.get("Type"): node.get("Target") for node in rels}
check_text("la relazione officeDocument punta al documento",
           targets.get(docs.REL_OFFICE_DOCUMENT), "word/document.xml")
check_true("...e il bersaglio esiste davvero nel pacchetto",
           all(target in names for target in targets.values()))

body = archive.read("word/document.xml").decode("utf-8")
document = ElementTree.fromstring(body)
w = "{" + docs.NS_WORD + "}"
paragraphs = document.iter(w + "p")
texts = [node.text or "" for node in document.iter(w + "t")]
print("        {0} paragrafi, {1} porzioni di testo".format(
    len(list(document.iter(w + "p"))), len(texts)))
check_true("il titolo e' nel documento", any(report.title in t for t in texts))
check_true("...e il nome del progetto",
           any("Podere della Fonte" in t for t in texts))
check_true("...e una voce dei vincoli", any("Strade" in t for t in texts))
rows = list(document.iter(w + "tr"))
check_true("le tabelle sono tabelle Word, non testo allineato",
           len(list(document.iter(w + "tbl"))) >= 3 and len(rows) > 6)
check_true("l'elenco delle piante non e' finito nel documento",
           not any("Progressivo" in t for t in texts))
check_true("il documento dichiara la pagina",
           len(list(document.iter(w + "sectPr"))) == 1)
core = ElementTree.fromstring(archive.read("docProps/core.xml"))
creators = [node.text for node in core if node.tag.endswith("creator")]
check_true("le proprieta' dicono chi l'ha redatta",
           creators and "Mancini" in (creators[0] or ""))
archive.close()

# --------------------------------------------------------------------------
# D5 - and Word itself, the only authority on a .docx
# --------------------------------------------------------------------------
print("\n== D5: lo apre Microsoft Word? ==")
word = None
try:
    import win32com.client                                      # noqa: E402

    word = win32com.client.DispatchEx("Word.Application")
except Exception as exc:                                        # noqa: BLE001
    skip("Microsoft Word apre il documento",
         "Word non disponibile su questa macchina ({0})".format(
             type(exc).__name__))

if word is not None:
    try:
        word.Visible = False
        word.DisplayAlerts = 0
        opened = word.Documents.Open(os.path.abspath(docx_path),
                                     ReadOnly=True, AddToRecentFiles=False)
        content = opened.Content.Text
        table_count = opened.Tables.Count
        paragraph_count = opened.Paragraphs.Count
        title_property = opened.BuiltInDocumentProperties("Title").Value
        opened.Close(False)
        print("        Word: {0:,} caratteri, {1} tabelle, {2} paragrafi"
              .format(len(content), table_count, paragraph_count))
        check_true("Word apre il documento e ne legge il testo",
                   len(content) > 200)
        check_true("...col titolo della relazione",
                   "RELAZIONE TECNICA" in content)
        check_true("...col nome del progetto",
                   "Podere della Fonte" in content)
        check_true("...con la voce dei vincoli", "Strade" in content)
        check_true("...e le tabelle sono tabelle per Word", table_count >= 3)
        check_true("le proprieta' del documento arrivano a Word",
                   report.title in str(title_property))
    except Exception as exc:                                    # noqa: BLE001
        check_true("Microsoft Word apre il documento senza errori: {0}"
                   .format(exc), False)
    finally:
        try:
            word.Quit()
        except Exception:                                       # noqa: BLE001
            pass

# --------------------------------------------------------------------------
# D6 - the three files say the same thing, and refusals stay refusals
# --------------------------------------------------------------------------
print("\n== D6: i tre file dicono la stessa cosa ==")
plain = docs.as_text(report)
html = docs.as_html(report)
for needle in ("Podere della Fonte", "Strade", "quercia"):
    check_true("'{0}' e' nel testo".format(needle), needle in plain)
    check_true("'{0}' e' nel PDF sorgente".format(needle), needle in html)
zone_row = tables["Zone"].rows[0]
check_true("la superficie della prima zona compare identica nei due",
           str(zone_row[1]) in plain and str(zone_row[1]) in html)

messages = []
outputs.warn = lambda exc: messages.append(exc.formatted())
check_text("un formato inventato non scrive niente",
           outputs.write_document("wpd", os.path.join(TMP, "x.wpd")), "")
check_true("...e lo dice", messages and "Formato" in messages[-1])
print("        rifiuto: {0}".format(messages[-1]))

check_text("un nome di foglio con caratteri vietati viene ripulito",
           docs.sheet_name("Zone/Fasce [2024]"), "Zone Fasce  2024")
check_text("...e non si sovrappone a uno gia' usato",
           docs.sheet_name("Zone", {"Zone"}), "Zone (2)")
check("...e non supera i 31 caratteri di Excel",
      len(docs.sheet_name("x" * 60)), 31)

workspace.unmount()

print("\n" + "=" * 78)
plan = report = book = prose = plants_sheet = zone_sheet = None
archive = document = types = rels = core = None
workspace = state = scheme = zones_panel = generate_panel = outputs = None
gc.collect()
QgsProject.instance().removeAllMapLayers()
QGS.exitQgis()
if SKIPS:
    print("SKIPPED ({0}):".format(len(SKIPS)))
    for label, reason in SKIPS:
        print("   - {0}: {1}".format(label, reason))
if FAILURES:
    print("FAILED ({0}):".format(len(FAILURES)))
    for name in FAILURES:
        print("   - {0}".format(name))
    sys.exit(1)
print("ALL CHECKS PASSED")
