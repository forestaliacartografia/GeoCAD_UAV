"""
The written deliverables: one report, three files.

A reforestation project is handed over as paper and as spreadsheets. The
numbers in the PDF, in the Word file and in the workbook have to be the same
numbers, so there is one report object here and three writers for it -- not
three functions that each go and ask the model again and drift apart by a
rounding.

``Report`` is a list of blocks: headings, paragraphs and tables. Nothing in
it knows about QGIS, so the whole thing can be built, inspected and compared
without a map.

**PDF** goes through Qt's own ``QTextDocument`` and ``QPdfWriter``. They ship
with PyQt, which the plugin already requires, so the report needs no library
the operator has to install.

**XLSX** goes through ``openpyxl``, which QGIS ships (3.40.15: 3.1.2, 4.0.0:
3.1.5). Verified by writing a workbook and reading it back with the same
library.

**DOCX** is written directly as an Open Packaging Conventions zip. QGIS ships
no Word library, and rather than refuse the format the package is composed by
hand from the parts Word itself writes -- the content types, the relationship
types and the WordprocessingML in this module were read off a .docx produced
by Microsoft Word 2016 and checked against Microsoft's own documentation of
the minimum document scenario (document / body / p / r / t). It is written
minimal on purpose: main document part, styles, core properties. What is
*not* claimed is any feature beyond text, headings and tables -- no headers,
no footers, no fields, and no images, because those were not verified.

**Figures** (1.35.0) are carried as PNG bytes and appear in the PDF and in
the HTML. The PDF path registers them on the ``QTextDocument`` as
``ResourceType.ImageResource`` -- verified present in its scoped spelling on
Qt 5.15 (QGIS 3.40.15) and Qt 6.8 (QGIS 4.0.0), with a document that
measurably lays the image out. The HTML embeds them as data URIs so the file
stays one file. The workbook and the Word document say where the figure is
rather than carrying it, for the reason above.
"""

from __future__ import annotations

import base64 as _b64
import datetime as _datetime
import os
import zipfile
from dataclasses import dataclass, field
from typing import Optional
from xml.sax.saxutils import escape

from ..core.errors import ExportError, InvalidInputError

# --------------------------------------------------------------------------
# The report, once
# --------------------------------------------------------------------------

BLOCK_HEADING = "heading"
BLOCK_TEXT = "text"
BLOCK_TABLE = "table"
BLOCK_IMAGE = "image"

#: Where a figure is referred to in the writers that cannot carry it.
IMAGE_ELSEWHERE = "(figura disponibile nelle versioni PDF e HTML)"


@dataclass
class Block:
    """One piece of the report. A heading, a paragraph or a table."""

    kind: str
    text: str = ""
    level: int = 1
    columns: tuple = ()
    rows: tuple = ()
    #: A table that belongs in a spreadsheet and nowhere else. Two thousand
    #: plants are a column in Excel and forty unreadable pages in a PDF, so
    #: the prose writers skip these and the workbook keeps them.
    sheet_only: bool = False
    #: PNG bytes, for a figure. The caption is ``text``.
    image: bytes = b""

    def is_empty(self) -> bool:
        if self.kind == BLOCK_TABLE:
            return not self.rows
        if self.kind == BLOCK_IMAGE:
            return not self.image
        return not str(self.text).strip()


@dataclass
class Report:
    """What the deliverables say, in one place."""

    title: str = "Relazione tecnica"
    subtitle: str = ""
    author: str = ""
    date: str = ""
    blocks: list = field(default_factory=list)

    def heading(self, text: str, level: int = 1) -> "Report":
        self.blocks.append(Block(BLOCK_HEADING, text=text, level=level))
        return self

    def text(self, text: str = "") -> "Report":
        self.blocks.append(Block(BLOCK_TEXT, text=text))
        return self

    def lines(self, lines) -> "Report":
        """A describe() list: its first line is the heading it always is."""
        items = [str(line) for line in lines if str(line).strip()]
        if not items:
            return self
        self.heading(items[0].strip(), level=2)
        for line in items[1:]:
            self.text(line.strip())
        return self

    def table(self, title: str, columns, rows,
              sheet_only: bool = False) -> "Report":
        kept = [tuple(row) for row in rows]
        self.blocks.append(Block(BLOCK_TABLE, text=title,
                                 columns=tuple(columns), rows=tuple(kept),
                                 sheet_only=sheet_only))
        return self

    def image_block(self, png: bytes, caption: str = "") -> "Report":
        """A figure, as PNG bytes, with the caption it is known by."""
        if png:
            self.blocks.append(Block(BLOCK_IMAGE, text=caption,
                                     image=bytes(png)))
        return self

    def images(self):
        return [block for block in self.blocks
                if block.kind == BLOCK_IMAGE and block.image]

    def prose(self):
        """The blocks the written report shows: everything but the listings."""
        return [block for block in self.blocks if not block.sheet_only]

    def tables(self):
        return [block for block in self.blocks if block.kind == BLOCK_TABLE]

    def stamp(self) -> str:
        parts = [part for part in (self.author, self.date) if part]
        return " - ".join(parts)


def as_text(report: Report) -> str:
    """The report as plain text -- what the panel has always shown."""
    lines = [report.title]
    if report.subtitle:
        lines.append(report.subtitle)
    if report.stamp():
        lines.append(report.stamp())
    lines.append("")
    for block in report.prose():
        if block.kind == BLOCK_HEADING:
            lines.append("")
            lines.append(block.text)
        elif block.kind == BLOCK_IMAGE:
            lines.append("")
            lines.append("[{0}] {1}".format(block.text or "figura",
                                            IMAGE_ELSEWHERE))
        elif block.kind == BLOCK_TABLE:
            lines.append("")
            lines.append(block.text)
            widths = _widths(block)
            lines.append("  " + _row_text(block.columns, widths))
            for row in block.rows:
                lines.append("  " + _row_text(row, widths))
        else:
            lines.append(block.text)
    return "\n".join(lines)


def _widths(block: Block):
    widths = [len(str(name)) for name in block.columns]
    for row in block.rows:
        for index, value in enumerate(row):
            if index < len(widths):
                widths[index] = max(widths[index], len(str(value)))
    return widths


def _row_text(values, widths) -> str:
    cells = []
    for index, value in enumerate(values):
        width = widths[index] if index < len(widths) else 0
        cells.append(str(value).ljust(width))
    return "  ".join(cells).rstrip()


def today() -> str:
    return _datetime.date.today().strftime("%d/%m/%Y")


# --------------------------------------------------------------------------
# PDF -- Qt's own text engine
# --------------------------------------------------------------------------

#: Page margins for the written report, in millimetres.
PAGE_MARGIN_MM = 18.0


def image_name(index: int) -> str:
    """The resource name a figure is registered and referred to by."""
    return "geocad-figura-{0}".format(int(index))


def as_html(report: Report, embed_images: bool = True) -> str:
    """The report as a document Qt can lay out. Also what the DOCX mirrors.

    ``embed_images`` decides how a figure is referred to: a ``data:`` URI,
    which keeps a standalone HTML file in one piece, or the resource name
    :func:`write_pdf` registers on the QTextDocument, which is the only form
    Qt's own text engine resolves.
    """
    out = ["<html><head><meta charset='utf-8'></head><body>",
           "<h1>{0}</h1>".format(escape(report.title))]
    if report.subtitle:
        out.append("<p><i>{0}</i></p>".format(escape(report.subtitle)))
    if report.stamp():
        out.append("<p><small>{0}</small></p>".format(escape(report.stamp())))
    seen_images = 0
    for block in report.prose():
        if block.kind == BLOCK_HEADING:
            level = 2 if block.level <= 1 else 3
            out.append("<h{0}>{1}</h{0}>".format(level, escape(block.text)))
        elif block.kind == BLOCK_IMAGE:
            if embed_images:
                source = "data:image/png;base64,{0}".format(
                    _b64.b64encode(block.image).decode("ascii"))
            else:
                source = image_name(seen_images)
            seen_images += 1
            out.append("<p><img src='{0}'></p>".format(source))
            if block.text:
                out.append("<p><small><i>{0}</i></small></p>".format(
                    escape(block.text)))
        elif block.kind == BLOCK_TABLE:
            out.append("<h3>{0}</h3>".format(escape(block.text)))
            out.append("<table border='1' cellspacing='0' cellpadding='3' "
                       "width='100%'>")
            out.append("<tr>" + "".join(
                "<th align='left'>{0}</th>".format(escape(str(name)))
                for name in block.columns) + "</tr>")
            for row in block.rows:
                out.append("<tr>" + "".join(
                    "<td>{0}</td>".format(escape(str(value)))
                    for value in row) + "</tr>")
            out.append("</table>")
        else:
            out.append("<p>{0}</p>".format(
                escape(block.text).replace("  ", "&nbsp;&nbsp;")))
    out.append("</body></html>")
    return "\n".join(out)


def write_pdf(report: Report, path: str) -> str:
    """Write the report to PDF with the text engine PyQt already carries."""
    from qgis.PyQt.QtCore import QMarginsF, QUrl                 # noqa: PLC0415
    from qgis.PyQt.QtGui import (QImage, QPageLayout,            # noqa: PLC0415
                                 QPageSize, QPdfWriter, QTextDocument)

    writer = QPdfWriter(path)
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageMargins(QMarginsF(PAGE_MARGIN_MM, PAGE_MARGIN_MM,
                                    PAGE_MARGIN_MM, PAGE_MARGIN_MM),
                          QPageLayout.Unit.Millimeter)
    writer.setTitle(report.title)
    document = QTextDocument()
    # Registered before setHtml: the layout resolves <img src> while it
    # parses, and a resource added afterwards arrives too late to be placed.
    for index, block in enumerate(report.images()):
        picture = QImage()
        if picture.loadFromData(block.image, "PNG"):
            document.addResource(
                QTextDocument.ResourceType.ImageResource,
                QUrl(image_name(index)), picture)
    document.setHtml(as_html(report, embed_images=False))
    # PyQt5 spells it print_, PyQt6 print. One plugin, two QGIS majors.
    render = getattr(document, "print_", None) or getattr(document, "print")
    render(writer)
    del writer
    _written(path, "PDF")
    return path


# --------------------------------------------------------------------------
# XLSX -- openpyxl, which QGIS ships
# --------------------------------------------------------------------------

#: Excel refuses these in a sheet name, and truncates past 31 characters.
SHEET_FORBIDDEN = ":\\/?*[]"
SHEET_MAX = 31


def sheet_name(title: str, used=()) -> str:
    """A sheet name Excel will accept, unique among the ones already used."""
    cleaned = "".join(" " if ch in SHEET_FORBIDDEN else ch
                      for ch in str(title)).strip() or "Foglio"
    cleaned = cleaned[:SHEET_MAX]
    if cleaned not in used:
        return cleaned
    for index in range(2, 100):
        suffix = " ({0})".format(index)
        candidate = cleaned[:SHEET_MAX - len(suffix)] + suffix
        if candidate not in used:
            return candidate
    return cleaned[:SHEET_MAX - 4] + " (99)"


def write_xlsx(report: Report, path: str) -> str:
    """One sheet for the prose, one for each table. Numbers stay numbers."""
    try:
        from openpyxl import Workbook                            # noqa: PLC0415
        from openpyxl.styles import Font                         # noqa: PLC0415
    except ImportError as exc:
        raise ExportError(
            "openpyxl is not available: {0}".format(exc),
            user_message="Questa installazione di QGIS non comprende "
                         "openpyxl: impossibile scrivere il foglio di "
                         "calcolo.",
            hint="Esporta in CSV oppure in PDF.") from exc

    book = Workbook()
    sheet = book.active
    sheet.title = "Relazione"
    bold = Font(bold=True)
    row_index = 1
    sheet.cell(row=row_index, column=1, value=report.title).font = bold
    row_index += 1
    for line in (report.subtitle, report.stamp()):
        if line:
            sheet.cell(row=row_index, column=1, value=line)
            row_index += 1
    row_index += 1
    for block in report.prose():
        if block.kind == BLOCK_HEADING:
            row_index += 1
            sheet.cell(row=row_index, column=1, value=block.text).font = bold
            row_index += 1
        elif block.kind == BLOCK_TEXT:
            sheet.cell(row=row_index, column=1, value=block.text)
            row_index += 1
        elif block.kind == BLOCK_IMAGE:
            sheet.cell(row=row_index, column=1,
                       value="{0} {1}".format(block.text or "Figura",
                                              IMAGE_ELSEWHERE))
            row_index += 1
    sheet.column_dimensions["A"].width = 70

    used = {sheet.title}
    for block in report.tables():
        name = sheet_name(block.text, used)
        used.add(name)
        page = book.create_sheet(name)
        for column, header in enumerate(block.columns, start=1):
            page.cell(row=1, column=column, value=str(header)).font = bold
        for line, row in enumerate(block.rows, start=2):
            for column, value in enumerate(row, start=1):
                page.cell(row=line, column=column, value=_cell(value))
        page.freeze_panes = "A2"
    book.save(path)
    _written(path, "foglio di calcolo")
    return path


def _cell(value):
    """Keep a number a number: a spreadsheet of strings cannot be summed."""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    stripped = text.replace(".", "", 1) if text.count(".") == 1 else text
    if stripped.lstrip("-").isdigit():
        try:
            return float(text) if "." in text else int(text)
        except ValueError:
            return text
    return text


# --------------------------------------------------------------------------
# DOCX -- the package, written by hand
# --------------------------------------------------------------------------

#: Namespaces and content types, read off a .docx written by Microsoft Word
#: 2016 and matched against Microsoft's documentation of the minimum
#: WordprocessingML document.
NS_CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
NS_RELATIONSHIPS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_WORD = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_CORE = ("http://schemas.openxmlformats.org/package/2006/"
           "metadata/core-properties")
NS_DC = "http://purl.org/dc/elements/1.1/"
NS_DCTERMS = "http://purl.org/dc/terms/"
NS_XSI = "http://www.w3.org/2001/XMLSchema-instance"

CT_DOCUMENT = ("application/vnd.openxmlformats-officedocument."
               "wordprocessingml.document.main+xml")
CT_STYLES = ("application/vnd.openxmlformats-officedocument."
             "wordprocessingml.styles+xml")
CT_CORE = "application/vnd.openxmlformats-package.core-properties+xml"
CT_RELS = "application/vnd.openxmlformats-package.relationships+xml"

REL_OFFICE_DOCUMENT = ("http://schemas.openxmlformats.org/officeDocument/"
                       "2006/relationships/officeDocument")
REL_STYLES = ("http://schemas.openxmlformats.org/officeDocument/2006/"
              "relationships/styles")
REL_CORE = ("http://schemas.openxmlformats.org/package/2006/relationships/"
            "metadata/core-properties")

#: A4 portrait in twentieths of a point, which is what WordprocessingML
#: measures pages in: 210 x 297 mm.
PAGE_W_TWIP = 11906
PAGE_H_TWIP = 16838
MARGIN_TWIP = 1134                              # 20 mm

XML_HEADER = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'


def _w(text: str) -> str:
    return escape(str(text))


def _paragraph(text: str, style: str = "") -> str:
    style_xml = ('<w:pPr><w:pStyle w:val="{0}"/></w:pPr>'.format(style)
                 if style else "")
    if not str(text).strip():
        return "<w:p>{0}</w:p>".format(style_xml)
    # xml:space preserve, or Word eats the leading spaces the describe()
    # lines use to indent their values.
    return ('<w:p>{0}<w:r><w:t xml:space="preserve">{1}</w:t></w:r></w:p>'
            .format(style_xml, _w(text)))


def _table(block: Block) -> str:
    columns = len(block.columns) or 1
    width = (PAGE_W_TWIP - 2 * MARGIN_TWIP) // columns
    out = ['<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/>'
           '<w:tblW w:w="0" w:type="auto"/>'
           '<w:tblBorders>'
           '<w:top w:val="single" w:sz="4" w:color="808080"/>'
           '<w:left w:val="single" w:sz="4" w:color="808080"/>'
           '<w:bottom w:val="single" w:sz="4" w:color="808080"/>'
           '<w:right w:val="single" w:sz="4" w:color="808080"/>'
           '<w:insideH w:val="single" w:sz="4" w:color="808080"/>'
           '<w:insideV w:val="single" w:sz="4" w:color="808080"/>'
           '</w:tblBorders></w:tblPr>']
    out.append("<w:tblGrid>" + "".join(
        '<w:gridCol w:w="{0}"/>'.format(width) for _ in range(columns))
        + "</w:tblGrid>")
    header = "".join(
        '<w:tc><w:tcPr><w:tcW w:w="{0}" w:type="dxa"/></w:tcPr>'
        '<w:p><w:r><w:rPr><w:b/></w:rPr><w:t xml:space="preserve">{1}</w:t>'
        '</w:r></w:p></w:tc>'.format(width, _w(name))
        for name in block.columns)
    out.append("<w:tr>{0}</w:tr>".format(header))
    for row in block.rows:
        cells = "".join(
            '<w:tc><w:tcPr><w:tcW w:w="{0}" w:type="dxa"/></w:tcPr>'
            '<w:p><w:r><w:t xml:space="preserve">{1}</w:t></w:r></w:p>'
            '</w:tc>'.format(width, _w(value)) for value in row)
        out.append("<w:tr>{0}</w:tr>".format(cells))
    out.append("</w:tbl>")
    out.append(_paragraph(""))
    return "".join(out)


def document_xml(report: Report) -> str:
    """The main document part: the story, and nothing else."""
    body = [_paragraph(report.title, "Title")]
    if report.subtitle:
        body.append(_paragraph(report.subtitle, "Subtitle"))
    if report.stamp():
        body.append(_paragraph(report.stamp()))
    for block in report.prose():
        if block.kind == BLOCK_HEADING:
            body.append(_paragraph(
                block.text, "Heading1" if block.level <= 1 else "Heading2"))
        elif block.kind == BLOCK_IMAGE:
            # A picture in WordprocessingML needs a media part, a
            # relationship and a drawing element, none of which was read off
            # a real .docx: the caption and where to find the figure, then.
            body.append(_paragraph("{0} {1}".format(
                block.text or "Figura", IMAGE_ELSEWHERE)))
        elif block.kind == BLOCK_TABLE:
            body.append(_paragraph(block.text, "Heading2"))
            body.append(_table(block))
        else:
            body.append(_paragraph(block.text))
    body.append(
        '<w:sectPr><w:pgSz w:w="{0}" w:h="{1}"/>'
        '<w:pgMar w:top="{2}" w:right="{2}" w:bottom="{2}" w:left="{2}"'
        ' w:header="708" w:footer="708" w:gutter="0"/></w:sectPr>'.format(
            PAGE_W_TWIP, PAGE_H_TWIP, MARGIN_TWIP))
    return (XML_HEADER
            + '<w:document xmlns:w="{0}"><w:body>{1}</w:body></w:document>'
            .format(NS_WORD, "".join(body)))


def styles_xml() -> str:
    """Just enough style for a heading to look like one."""
    def style(style_id, name, size_half_points, bold, outline):
        return (
            '<w:style w:type="paragraph" w:styleId="{0}">'
            '<w:name w:val="{1}"/><w:basedOn w:val="Normal"/>'
            '<w:pPr><w:outlineLvl w:val="{2}"/>'
            '<w:spacing w:before="200" w:after="80"/></w:pPr>'
            '<w:rPr>{3}<w:sz w:val="{4}"/><w:szCs w:val="{4}"/></w:rPr>'
            '</w:style>'.format(style_id, name, outline,
                                "<w:b/>" if bold else "", size_half_points))

    return (XML_HEADER + '<w:styles xmlns:w="{0}">'.format(NS_WORD)
            + '<w:docDefaults><w:rPrDefault><w:rPr>'
              '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
              '<w:sz w:val="20"/><w:szCs w:val="20"/>'
              '</w:rPr></w:rPrDefault></w:docDefaults>'
            + '<w:style w:type="paragraph" w:default="1" w:styleId="Normal">'
              '<w:name w:val="Normal"/></w:style>'
            + style("Title", "Title", 40, True, 0)
            + style("Subtitle", "Subtitle", 24, False, 1)
            + style("Heading1", "heading 1", 28, True, 0)
            + style("Heading2", "heading 2", 24, True, 1)
            + '<w:style w:type="table" w:styleId="TableGrid">'
              '<w:name w:val="Table Grid"/></w:style>'
            + '</w:styles>')


def core_xml(report: Report) -> str:
    stamp = _datetime.datetime.now(_datetime.timezone.utc).strftime(
        "%Y-%m-%dT%H:%M:%SZ")
    return (XML_HEADER
            + '<cp:coreProperties xmlns:cp="{0}" xmlns:dc="{1}" '
              'xmlns:dcterms="{2}" xmlns:xsi="{3}">'
              '<dc:title>{4}</dc:title><dc:creator>{5}</dc:creator>'
              '<cp:lastModifiedBy>{5}</cp:lastModifiedBy>'
              '<dcterms:created xsi:type="dcterms:W3CDTF">{6}</dcterms:created>'
              '<dcterms:modified xsi:type="dcterms:W3CDTF">{6}</dcterms:modified>'
              '</cp:coreProperties>'.format(
                  NS_CORE, NS_DC, NS_DCTERMS, NS_XSI, _w(report.title),
                  _w(report.author or "GeoCad UAV Toolkit"), stamp))


def content_types_xml() -> str:
    return (XML_HEADER
            + '<Types xmlns="{0}">'
              '<Default Extension="rels" ContentType="{1}"/>'
              '<Default Extension="xml" ContentType="application/xml"/>'
              '<Override PartName="/word/document.xml" ContentType="{2}"/>'
              '<Override PartName="/word/styles.xml" ContentType="{3}"/>'
              '<Override PartName="/docProps/core.xml" ContentType="{4}"/>'
              '</Types>'.format(NS_CONTENT_TYPES, CT_RELS, CT_DOCUMENT,
                                CT_STYLES, CT_CORE))


def package_rels_xml() -> str:
    return (XML_HEADER
            + '<Relationships xmlns="{0}">'
              '<Relationship Id="rId1" Type="{1}" Target="word/document.xml"/>'
              '<Relationship Id="rId2" Type="{2}" Target="docProps/core.xml"/>'
              '</Relationships>'.format(NS_RELATIONSHIPS,
                                        REL_OFFICE_DOCUMENT, REL_CORE))


def document_rels_xml() -> str:
    return (XML_HEADER
            + '<Relationships xmlns="{0}">'
              '<Relationship Id="rId1" Type="{1}" Target="styles.xml"/>'
              '</Relationships>'.format(NS_RELATIONSHIPS, REL_STYLES))


def write_docx(report: Report, path: str) -> str:
    """Write the report as an Open Packaging Conventions .docx."""
    parts = {
        "[Content_Types].xml": content_types_xml(),
        "_rels/.rels": package_rels_xml(),
        "word/document.xml": document_xml(report),
        "word/_rels/document.xml.rels": document_rels_xml(),
        "word/styles.xml": styles_xml(),
        "docProps/core.xml": core_xml(report),
    }
    try:
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            # [Content_Types].xml first: some readers look for it at the
            # start of the archive rather than in the central directory.
            archive.writestr("[Content_Types].xml",
                             parts.pop("[Content_Types].xml"))
            for name, content in parts.items():
                archive.writestr(name, content)
    except OSError as exc:
        raise ExportError(
            "cannot write {0}: {1}".format(path, exc),
            user_message="Impossibile scrivere il documento Word.",
            hint=str(exc)) from exc
    _written(path, "documento Word")
    return path


# --------------------------------------------------------------------------

FORMATS = (
    ("pdf", "PDF (.pdf)", ".pdf", write_pdf),
    ("docx", "Word (.docx)", ".docx", write_docx),
    ("xlsx", "Excel (.xlsx)", ".xlsx", write_xlsx),
)


def writer_for(key: str):
    for candidate, _label, _suffix, function in FORMATS:
        if candidate == key:
            return function
    raise InvalidInputError(
        "no writer for {0!r}".format(key),
        user_message="Formato di relazione non riconosciuto.",
        hint="Usa: {0}.".format(", ".join(
            candidate for candidate, _l, _s, _w in FORMATS)))


def suffix_for(key: str) -> str:
    for candidate, _label, suffix, _function in FORMATS:
        if candidate == key:
            return suffix
    return ""


def write(report: Report, path: str, key: str) -> str:
    return writer_for(key)(report, path)


def _written(path: str, what: str) -> None:
    if not os.path.exists(path) or os.path.getsize(path) <= 0:
        raise ExportError(
            "{0} produced nothing at {1}".format(what, path),
            user_message="L'esportazione in {0} non ha prodotto alcun "
                         "file.".format(what))


__all__ = ["Report", "Block", "BLOCK_HEADING", "BLOCK_TEXT", "BLOCK_TABLE",
           "BLOCK_IMAGE", "IMAGE_ELSEWHERE", "image_name",
           "FORMATS", "as_text", "as_html", "write", "write_pdf",
           "write_docx", "write_xlsx", "writer_for", "suffix_for",
           "sheet_name", "today", "document_xml", "content_types_xml",
           "package_rels_xml", "document_rels_xml", "styles_xml", "core_xml"]
